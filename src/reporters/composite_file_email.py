"""
Модуль для создания комплексных отчетов и отправки по email.

Генерирует отчеты в различных форматах (PNG, HTML, PDF, Excel) и отправляет их по email.
"""

import os
import smtplib
from email import encoders
from email.mime.base import MIMEBase
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText

import matplotlib.pyplot as plt
import numpy as np
import openpyxl
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import seaborn as sns
from loguru import logger
from openpyxl.chart import BarChart, Reference
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer

from src.config import AppConfig
from src.config.connections import SMTPConnection
from src.config.pipelines import Pipeline
from src.reporters.base import BaseReporter


class CompositeFileEmailReporter(BaseReporter):
    """Репортер для создания комплексных отчетов и отправки по email."""

    def generate_reports(
        self, pipeline: Pipeline, df: pd.DataFrame, df_no_scale: pd.DataFrame
    ):
        """Генерирует все типы отчетов для пайплайна."""
        self.generated_files = []
        plt.style.use("seaborn-v0_8")
        plt.rcParams["font.size"] = 10

        self._generate_static_plots(pipeline, df, df_no_scale)
        self._generate_interactive_plots(pipeline, df, df_no_scale)
        self._generate_pdf_report(pipeline)
        self._generate_excel_report(pipeline)

    def _generate_static_plots(self, pipeline, df, df_no_scale):
        """Создает статические графики matplotlib."""
        numeric_cols = df_no_scale.select_dtypes(include=[np.number]).columns
        target_col = pipeline.model.target

        # Корреляционная матрица
        plt.figure(figsize=(12, 10))
        show_annotations = len(numeric_cols) <= 20
        sns.heatmap(
            df_no_scale[numeric_cols].corr(),
            annot=show_annotations,
            cmap="coolwarm",
            center=0,
            fmt=".2f",
            cbar_kws={"shrink": 0.8},
        )
        plt.title("Матрица корреляции признаков", fontsize=14, pad=20)
        plt.tight_layout()
        corr_path = os.path.join(pipeline.meta.results_path, "correlation_heatmap.png")
        plt.savefig(corr_path, dpi=300, bbox_inches="tight")
        plt.close()
        self.generated_files.append(corr_path)

        # Распределение целевой переменной
        plt.figure(figsize=(10, 6))
        if pipeline.model.type == "classification":
            counts = df_no_scale[target_col].value_counts()
            ax = counts.plot(kind="bar", alpha=0.8, color="skyblue", edgecolor="black")
            plt.title(f"Распределение целевой переменной: {target_col}", fontsize=14)
            plt.xlabel("Классы")
            plt.ylabel("Количество")
            plt.xticks(rotation=0)
            for i, v in enumerate(counts.values):
                ax.text(
                    i, v + max(counts.values) * 0.01, str(v), ha="center", va="bottom"
                )
        else:
            plt.hist(
                df_no_scale[target_col],
                bins=30,
                alpha=0.7,
                color="lightgreen",
                edgecolor="black",
            )
            plt.title(f"Распределение целевой переменной: {target_col}", fontsize=14)
            plt.xlabel(target_col)
            plt.ylabel("Частота")
        plt.tight_layout()
        target_path = os.path.join(
            pipeline.meta.results_path, "target_distribution.png"
        )
        plt.savefig(target_path, dpi=300, bbox_inches="tight")
        plt.close()
        self.generated_files.append(target_path)

        # Метрики модели
        plt.figure(figsize=(10, 6))
        metrics_names = list(pipeline.meta.metrics.keys())
        metrics_values = list(pipeline.meta.metrics.values())

        min_val, max_val = min(metrics_values), max(metrics_values)
        y_range = max_val - min_val
        if y_range < 0.1:
            y_bottom = max(0, min_val - 0.05)
            y_top = min(1, max_val + 0.05)
            plt.ylim(y_bottom, y_top)

        bars = plt.bar(
            metrics_names, metrics_values, alpha=0.8, color="coral", edgecolor="black"
        )
        plt.title("Метрики производительности модели", fontsize=14)
        plt.ylabel("Значение")
        plt.xticks(rotation=45, ha="right")

        for bar, value in zip(bars, metrics_values):
            plt.text(
                bar.get_x() + bar.get_width() / 2,
                bar.get_height() + y_range * 0.01,
                f"{value:.3f}",
                ha="center",
                va="bottom",
                fontweight="bold",
            )

        plt.tight_layout()
        metrics_path = os.path.join(pipeline.meta.results_path, "model_metrics.png")
        plt.savefig(metrics_path, dpi=300, bbox_inches="tight")
        plt.close()
        self.generated_files.append(metrics_path)

        # Важность признаков
        if hasattr(pipeline.meta.trained_model, "feature_importances_"):
            importances = pipeline.meta.trained_model.feature_importances_
            feature_names = df.drop(columns=[target_col]).columns[: len(importances)]

            plt.figure(figsize=(10, 8))
            indices = np.argsort(importances)[::-1][:15]

            plt.barh(
                range(len(indices)),
                importances[indices],
                alpha=0.8,
                color="gold",
                edgecolor="black",
            )
            plt.yticks(range(len(indices)), [feature_names[i] for i in indices])
            plt.title("Важность признаков (топ-15)", fontsize=14)
            plt.xlabel("Важность")
            plt.gca().invert_yaxis()
            plt.tight_layout()

            importance_path = os.path.join(
                pipeline.meta.results_path, "feature_importance.png"
            )
            plt.savefig(importance_path, dpi=300, bbox_inches="tight")
            plt.close()
            self.generated_files.append(importance_path)

        # Компоненты временного ряда
        if pipeline.meta.is_timeseries and pipeline.time_column:
            plt.figure(figsize=(14, 8))

            plt.subplot(2, 1, 1)
            plt.plot(pipeline.meta.timeseries_stats.trend, color="blue", linewidth=2)
            plt.title("Тренд временного ряда", fontsize=12)
            plt.grid(True, alpha=0.3)

            plt.subplot(2, 1, 2)
            plt.plot(
                pipeline.meta.timeseries_stats.seasonal, color="green", linewidth=2
            )
            plt.title("Сезонная компонента", fontsize=12)
            plt.grid(True, alpha=0.3)

            plt.tight_layout()
            ts_path = os.path.join(
                pipeline.meta.results_path, "timeseries_components.png"
            )
            plt.savefig(ts_path, dpi=300, bbox_inches="tight")
            plt.close()
            self.generated_files.append(ts_path)

    def _generate_interactive_plots(self, pipeline, df, df_no_scale):
        """Создает интерактивные графики Plotly."""
        numeric_cols = df_no_scale.select_dtypes(include=[np.number]).columns
        target_col = pipeline.model.target

        # Интерактивные графики Plotly
        if len(numeric_cols) >= 2:
            fig = px.scatter(
                df_no_scale,
                x=numeric_cols[0],
                y=numeric_cols[1],
                color=target_col if pipeline.model.type == "classification" else None,
                title=f"Интерактивный график: {numeric_cols[0]} vs {numeric_cols[1]}",
            )
            plotly_scatter_path = os.path.join(
                pipeline.meta.results_path, "interactive_scatter.html"
            )
            fig.write_html(plotly_scatter_path)
            self.generated_files.append(plotly_scatter_path)

        metrics_names = list(pipeline.meta.metrics.keys())
        metrics_values = list(pipeline.meta.metrics.values())
        fig = go.Figure(
            data=[
                go.Bar(
                    x=metrics_names,
                    y=metrics_values,
                    text=[f"{v:.3f}" for v in metrics_values],
                    textposition="auto",
                    marker_color="lightcoral",
                )
            ]
        )
        fig.update_layout(
            title="Интерактивные метрики модели",
            xaxis_title="Метрики",
            yaxis_title="Значение",
        )
        plotly_metrics_path = os.path.join(
            pipeline.meta.results_path, "interactive_metrics.html"
        )
        fig.write_html(plotly_metrics_path)
        self.generated_files.append(plotly_metrics_path)

    def _generate_pdf_report(self, pipeline):
        """Создает PDF отчет с результатами анализа."""
        pdf_path = os.path.join(pipeline.meta.results_path, "model_report.pdf")
        doc = SimpleDocTemplate(pdf_path, pagesize=letter)

        pdfmetrics.registerFont(
            TTFont("DejaVuSerif", "dejavuserif/DejaVuSerif.ttf", "UTF-8")
        )
        styles = getSampleStyleSheet()
        styles["Normal"].fontName = "DejaVuSerif"
        styles["Title"].fontName = "DejaVuSerif"
        styles["Heading2"].fontName = "DejaVuSerif"

        story = []
        story.append(
            Paragraph(f"Отчет по ML пайплайну: {pipeline.name}", styles["Title"])
        )
        story.append(Spacer(1, 12))

        story.append(Paragraph(f"Тип модели: {pipeline.model.type}", styles["Normal"]))
        story.append(
            Paragraph(f"Целевая переменная: {pipeline.model.target}", styles["Normal"])
        )
        story.append(
            Paragraph(
                f"Размер тестовой выборки: {pipeline.model.test_size}", styles["Normal"]
            )
        )
        story.append(Spacer(1, 12))

        story.append(Paragraph("Метрики производительности:", styles["Heading2"]))
        for metric, value in pipeline.meta.metrics.items():
            story.append(Paragraph(f"{metric.upper()}: {value:.4f}", styles["Normal"]))
        story.append(Spacer(1, 12))

        # Добавляем изображения
        image_files = [f for f in self.generated_files if f.endswith(".png")]
        for img_path in image_files[:4]:
            if os.path.exists(img_path):
                story.append(Image(img_path, width=6 * inch, height=4 * inch))
                story.append(Spacer(1, 6))

        doc.build(story)
        self.generated_files.append(pdf_path)

    def _generate_excel_report(self, pipeline):
        """Создает Excel отчет с метриками и статистиками."""
        excel_path = os.path.join(pipeline.meta.results_path, "model_report.xlsx")
        workbook = openpyxl.Workbook()

        ws_metrics = workbook.active
        ws_metrics.title = "Метрики"
        ws_metrics["A1"] = "Метрика"
        ws_metrics["B1"] = "Значение"

        for i, (metric, value) in enumerate(pipeline.meta.metrics.items(), 2):
            ws_metrics[f"A{i}"] = metric.upper()
            ws_metrics[f"B{i}"] = round(value, 4)

        chart = BarChart()
        chart.title = "Метрики модели"
        data = Reference(
            ws_metrics, min_col=2, min_row=1, max_row=len(pipeline.meta.metrics) + 1
        )
        cats = Reference(
            ws_metrics, min_col=1, min_row=2, max_row=len(pipeline.meta.metrics) + 1
        )
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        ws_metrics.add_chart(chart, "D2")

        ws_stats = workbook.create_sheet("Статистики")
        ws_stats["A1"] = "Признак"
        ws_stats["B1"] = "Среднее"
        ws_stats["C1"] = "Медиана"
        ws_stats["D1"] = "Ст.отклонение"

        for i, col in enumerate(pipeline.meta.stats.mean.index, 2):
            ws_stats[f"A{i}"] = col
            ws_stats[f"B{i}"] = round(pipeline.meta.stats.mean[col], 4)
            ws_stats[f"C{i}"] = round(pipeline.meta.stats.median[col], 4)
            ws_stats[f"D{i}"] = round(pipeline.meta.stats.std[col], 4)

        workbook.save(excel_path)
        self.generated_files.append(excel_path)

    def send_email(self, pipeline: Pipeline, smtp_connection: SMTPConnection):
        """Отправляет отчет по email с вложениями."""
        msg = MIMEMultipart()
        msg["From"] = smtp_connection.username
        msg["To"] = ", ".join(pipeline.report.email_recipients)
        msg["Subject"] = f"Отчет по ML пайплайну: {pipeline.name}"
        msg.attach(
            MIMEText(
                f"Во вложении отчет по ML пайплайну: {pipeline.name}.\n\nС уважением,\nВаш smol-data-automation",
                "plain",
            )
        )

        allowed_extensions = (".png", ".pdf", ".xls", ".xlsx")
        for file_path in self.generated_files:
            if file_path.lower().endswith(allowed_extensions) and os.path.exists(
                file_path
            ):
                with open(file_path, "rb") as attachment:
                    part = MIMEBase("application", "octet-stream")
                    part.set_payload(attachment.read())
                    encoders.encode_base64(part)
                    part.add_header(
                        "Content-Disposition",
                        f"attachment; filename= {os.path.basename(file_path)}",
                    )
                    msg.attach(part)

        try:
            server = smtplib.SMTP_SSL(
                smtp_connection.host, smtp_connection.port, timeout=30
            )
        except Exception:
            server = smtplib.SMTP(
                smtp_connection.host, smtp_connection.port, timeout=30
            )
            server.starttls()

        server.login(smtp_connection.username, smtp_connection.password)
        server.send_message(msg)
        server.quit()

    async def report(
        self,
        pipeline: Pipeline,
        df: pd.DataFrame,
        df_no_scale: pd.DataFrame,
        config: AppConfig,
    ):
        """Основной метод для создания отчетов и отправки по email."""
        logger.info("Запускаю генерацию отчётов...")
        self.generate_reports(pipeline, df, df_no_scale)

        logger.info("Сгенерированы файлы отчётов:")
        for generated_file in self.generated_files:
            logger.info(generated_file)

        if not pipeline.report.smtp_connection_name:
            logger.info(
                "Отчёты не будут отправлены по email т.к. не указано SMTP-соединение"
            )
            return
        elif not pipeline.report.email_recipients:
            logger.info(
                "Отчёты не будут отправлены по email т.к. не указаны получатели"
            )
            return

        smtp_connection = config.get_smtp_connection(
            pipeline.report.smtp_connection_name
        )
        try:
            self.send_email(pipeline, smtp_connection)
        except Exception as e:
            logger.error(f"Email не удалось отправить из-за ошибки: {e}")
